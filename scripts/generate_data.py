import re, json, datetime, unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'source.xlsx'
OUT = ROOT / 'data.js'

STORY_ORDER = [
    '떡잎마을','진실호수','201번도로','잔모래마을','202번도로','축복시티','축복시티 트레이너 스쿨',
    '203번도로','무쇠게이트','무쇠시티','무쇠박물관','207번도로','204번도로','험한샛길','꽃향기마을','골짜기발전소','골풀무제철소','205번도로','영원의 숲','숲의 양옥집','영원시티','갤럭시단 영원 빌딩','206번도로','미혹의 동굴','천관산','208번도로','연고시티','209번도로','로스트타워','신수마을','신수유적','210번도로','봉신마을','215번도로','장막시티','214번도로','입지호수근처','입지호수','입지호수의 공동','213번도로','들판시티','212번도로','포켓몬저택','자랑의 뒷마당','218번도로','운하시티','강철섬','진실호수의 공동','216번도로','217번도로','예지호수근처','예지호수','예지호수의 공동','선단시티','선단신전','갤럭시단 장막 빌딩','갤럭시단 본부','창기둥','깨어진세계','222번도로','물가시티','223번수로','챔피언로드','포켓몬리그','224번도로','만월섬','신월섬','시작의 방','꽃의 낙원','배틀존','파이트에리어','225번도로','서바이벌에리어','226번수로','227번도로','하드마운틴','228번도로','229번도로','리조트에리어','230번수로','NPC 교환','특수 이벤트'
]

VALID_TYPES = {'노말','불꽃','물','풀','전기','얼음','격투','독','땅','비행','에스퍼','벌레','바위','고스트','드래곤','악','강철','페어리'}
REF_RE = re.compile(r'#\s*(\d{1,4})\s*([^\s\n#\(\),，:：\[\]]+)')

METHOD_PREFIXES = ('풀','동굴','파도타기','낚싯대','포켓트레','달콤한꿀','고정','선물','조우 방법','마박사','카우걸에게 받은 알','땅','물','바위깨기','낡은낚싯대','좋은낚싯대','대단한낚싯대','무리','대량발생','로밍','락클라임')
LOCATION_KEYWORDS = ('번도로','번수로','시티','마을','호수','숲','게이트','동굴','산','섬','유적','신전','저택','발전소','제철소','에리어','로드','낙원','세계','샛길','타워','고개','리그','사파리','리조트','하드마운틴','챔피언','방','본부','빌딩','트레이너 스쿨')


def clean(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    if isinstance(v, datetime.datetime):
        if v.year == 2025:
            a, b = sorted([v.month, v.day])
            return f'{a}-{b}'
        return v.strftime('%Y-%m-%d')
    s = str(v).replace('\n', ' ').strip()
    s = ''.join(ch for ch in s if unicodedata.category(ch)[0] != 'C')
    s = re.sub(r'\s+', ' ', s)
    s = s.replace('：', ':').replace('레벨.', '레벨:').replace('위치 :', '위치:')
    s = s.replace('(!)', '').strip()
    s = re.sub(r'\s+([,.])', r'\1', s)
    if s == '#NAME?':
        return ''
    if s.endswith('.0') and re.match(r'^-?\d+\.0$', s):
        s = s[:-2]
    return s


def clean_type(v):
    s = clean(v)
    s = re.sub(r'[^가-힣A-Za-z0-9]', '', s)
    return s if s in VALID_TYPES else ''


def parse_ref(v):
    s = clean(v)
    m = REF_RE.search(s)
    if not m:
        return None
    no = int(m.group(1))
    name = m.group(2).strip()
    if no < 1 or no > 493 or not name:
        return None
    return no, name


def number(ws, r, c):
    try:
        v = ws.cell(r, c).value
        if v in ('', None): return None
        return int(float(v))
    except Exception:
        return None


def looks_method(s):
    s = clean(s)
    if not s: return False
    return s.startswith(METHOD_PREFIXES) or any(k in s for k in ['조우','획득처','낚싯대','포켓트레','상세 조건'])


def looks_location(s):
    s = clean(s)
    if not s or s.startswith('-') or s.startswith('#') or '출현 포켓몬' in s:
        return False
    if re.match(r'^위치\s*:', s):
        return True
    if s in ('엔딩 후','깨어진세계','꽃의 낙원','시작의 방','배틀존','NPC 교환','특수 이벤트'):
        return True
    return any(k in s for k in LOCATION_KEYWORDS)


def normalize_location(s):
    s = clean(s)
    s = re.sub(r'^위치\s*:\s*', '', s).strip()
    s = s.replace(' ~ ', ' ').strip()
    if s == '트레이너 스쿨':
        return '축복시티 트레이너 스쿨'
    return s or '특수 이벤트'


def base_location(s):
    s = normalize_location(s)
    s = re.sub(r'\s*\([^)]*\)', '', s)
    s = re.sub(r'\s*[~/].*$', '', s)
    s = re.sub(r'\.\s*어디서든.*$', '', s)
    return s.strip()


def story_rank(name):
    b = base_location(name)
    for i, key in enumerate(STORY_ORDER):
        if b == key or b.startswith(key) or key in b or key in name:
            return i
    return 9999


def sprite_url(no):
    return f'https://raw.githubusercontent.com/PokeAPI/sprites/master/sprites/pokemon/versions/generation-iv/diamond-pearl/{int(no)}.png'


def add_unique(arr, e):
    def norm_note(n): return clean(n).rstrip('.')
    key = (e['pokemonId'], e['location'], e['method'], e.get('level',''), tuple(norm_note(x) for x in e.get('notes') or []), tuple((r.get('label',''),r.get('value','')) for r in e.get('rates') or []))
    if not hasattr(add_unique, 'seen'):
        add_unique.seen = set()
    if key not in add_unique.seen:
        add_unique.seen.add(key)
        arr.append(e)



def parse_evolution_links(no, evolution_text):
    """Return (pre_ids, evolves_to_ids, is_final) from Korean evolution text.

    Examples:
    - #004 레벨 16 - 진화 가능: #006  -> pre #004, evolves_to #006
    - #005 레벨 36                    -> pre #005, final
    - #002 로 진화 가능               -> evolves_to #002
    """
    text = clean(evolution_text)
    refs_all = [int(x) for x in re.findall(r'#\s*(\d{1,4})', text) if 1 <= int(x) <= 493]
    if not text or text in ('N', '-'):
        return [], [], True
    if '진화 가능' in text:
        before, after = text.split('진화 가능', 1)
        targets = [int(x) for x in re.findall(r'#\s*(\d{1,4})', after) if 1 <= int(x) <= 493 and int(x) != no]
        if targets:
            pre = [int(x) for x in re.findall(r'#\s*(\d{1,4})', before) if 1 <= int(x) <= 493 and int(x) != no]
        else:
            # Pattern such as "#002 로 진화 가능" keeps the target before the phrase.
            targets = [x for x in refs_all if x != no]
            pre = []
        return pre, targets, False
    # No forward-evolution phrase means this row is the final/terminal member for that branch.
    pre = [x for x in refs_all if x != no]
    return pre, [], True

def main():
    wb = load_workbook(SOURCE, data_only=True)
    fallback_images = {}
    if 'data' in wb.sheetnames:
        ws = wb['data']
        for r in range(1, ws.max_row + 1):
            try:
                no = int(float(ws.cell(r, 1).value))
            except Exception:
                continue
            url = clean(ws.cell(r, 9).value)
            if no and url.startswith('http'):
                fallback_images[no] = url

    ws = wb['STATS']
    pokemons = []
    seen_ids = set()
    current_gen = ''
    evolves_to = {}
    pre_ids = {}
    for r in range(4, ws.max_row + 1):
        gen = clean(ws.cell(r, 1).value)
        if gen:
            current_gen = gen
        ref = parse_ref(ws.cell(r, 2).value)
        if not ref:
            continue
        no, name = ref
        if no in seen_ids:
            continue
        seen_ids.add(no)
        hp, atk, defn, spa, spd, spe, total = [number(ws, r, c) for c in range(19, 26)]
        if hp is None:
            hp, atk, defn, spa, spd, spe, total = [number(ws, r, c) for c in range(5, 12)]
        types = []
        for c in (36, 38, 32, 34, 28, 30):
            t = clean_type(ws.cell(r,c).value)
            if t and t not in types:
                types.append(t)
            if len(types) >= 2:
                break
        # 레니게이드 플레티넘 컴플리트 기준 특성만 사용합니다.
        # 원본/클래식 특성까지 합치면 리자몽에 선파워처럼 실제 컴플리트 기준이 아닌 특성이 섞입니다.
        abilities = []
        for c in (44, 45):
            a = clean(ws.cell(r,c).value)
            a = re.sub(r'\s*\(!\)\s*', '', a).strip()
            if a and a not in abilities:
                abilities.append(a)
        evolution = clean(ws.cell(r, 68).value)
        pre, targets, is_final = parse_evolution_links(no, evolution)
        if targets:
            evolves_to[no] = targets
        if pre:
            pre_ids[no] = pre
        pokemons.append({
            'id': no, 'num': f'#{no:03d}', 'name': name, 'generation': current_gen,
            'sprite': sprite_url(no), 'image': fallback_images.get(no, ''),
            'types': types, 'abilities': abilities,
            'stats': {'hp': hp, 'atk': atk, 'def': defn, 'spa': spa, 'spd': spd, 'spe': spe, 'total': total},
            'evYield': clean(ws.cell(r, 52).value), 'expGrowth': clean(ws.cell(r, 56).value),
            'catchRate': clean(ws.cell(r, 58).value),
            'eggGroups': [x for x in [clean(ws.cell(r, 65).value), clean(ws.cell(r, 66).value)] if x],
            'evolution': evolution, 'evolvesTo': evolves_to.get(no, []), 'preEvolutionIds': pre_ids.get(no, []),
            'color': clean(ws.cell(r, 70).value), 'isFinalEvolution': is_final, 'locations': []
        })
    # patch pre/evolve maps after all rows
    pokemon_by_id = {p['id']: p for p in pokemons}
    for p in pokemons:
        if p['id'] in evolves_to:
            p['evolvesTo'] = [x for x in evolves_to[p['id']] if x in pokemon_by_id]
        if p['id'] in pre_ids:
            p['preEvolutionIds'] = [x for x in pre_ids[p['id']] if x in pokemon_by_id]
    for src, tgts in evolves_to.items():
        for target in tgts:
            if target in pokemon_by_id:
                pokemon_by_id[target].setdefault('preEvolutionIds', [])
                if src not in pokemon_by_id[target]['preEvolutionIds']:
                    pokemon_by_id[target]['preEvolutionIds'].append(src)

    encounters = []
    ws = wb['ENCOUNTERS']
    for c in range(5, ws.max_column + 1, 7):
        current_location = ''
        current_method = ''
        current_headers = ['아침','낮','밤']
        initial = clean(ws.cell(2,c).value)
        if looks_location(initial):
            current_location = normalize_location(initial)
        for r in range(2, ws.max_row + 1):
            first = clean(ws.cell(r,c).value)
            ref = parse_ref(ws.cell(r,c+1).value)
            if first and not parse_ref(first) and not ref:
                if looks_method(first):
                    current_method = first
                    h = [clean(ws.cell(r, c+i).value) for i in range(3,6)]
                    mapper = {'M':'아침','D':'낮','N':'밤','O':'낡은낚싯대','G':'좋은낚싯대','S':'대단한낚싯대'}
                    h = [mapper.get(x,x) for x in h if x]
                    current_headers = h if h else ['아침','낮','밤']
                elif looks_location(first):
                    current_location = normalize_location(first)
                    current_method = ''
                    current_headers = ['아침','낮','밤']
                continue
            if ref:
                no, _ = ref
                if no not in pokemon_by_id: continue
                level = clean(ws.cell(r,c+2).value)
                rates = []
                for i, head in zip(range(3,6), current_headers[:3]):
                    rv = clean(ws.cell(r,c+i).value)
                    if rv: rates.append({'label': head, 'value': rv})
                add_unique(encounters, {
                    'pokemonId': no, 'pokemonName': pokemon_by_id[no]['name'],
                    'location': current_location or '미분류 지역', 'method': ('상세 조건 확인' if str(current_method).startswith('조우 방법') else (current_method or '야생 출현')),
                    'level': level, 'rates': rates, 'source':'ENCOUNTERS', 'notes': []
                })

    gifts = []
    if 'GiFTSEVENTS' in wb.sheetnames:
        ws = wb['GiFTSEVENTS']
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell_text = clean(ws.cell(r,c).value)
                if not cell_text.startswith('#'):
                    continue
                ref = parse_ref(cell_text)
                if not ref: continue
                no, _ = ref
                if no not in pokemon_by_id: continue
                # Skip trainer-party blocks: they usually have a trainer label with [id] nearby and no location.
                loc_candidates, levels, notes, labels = [], [], [], []
                # Gift/event blocks are laid out in narrow vertical columns. Keep the search tight
                # so the next event block or special trainer table is not accidentally merged.
                for rr in range(max(1,r-6), min(ws.max_row,r+18)+1):
                    for cc in range(max(1,c-3), min(ws.max_column,c+3)+1):
                        t = clean(ws.cell(rr,cc).value)
                        if not t: continue
                        dist = abs(rr-r)+abs(cc-c)
                        if re.match(r'^위치\s*:', t):
                            loc_candidates.append((dist, normalize_location(t)))
                        elif re.match(r'^레벨\s*:', t):
                            levels.append((dist, re.sub(r'^레벨\s*:\s*','',t).strip()))
                        elif t.startswith('-'):
                            note = t.lstrip('-').strip()
                            # Prefer notes in same event column area, but include direct pokemon mentions.
                            if abs(cc-c) <= 2 or f'#{no:03d}' in note or pokemon_by_id[no]['name'] in note:
                                note = re.sub(r'\s+', ' ', note).strip()
                                if note and note not in notes:
                                    notes.append(note)
                        elif '선물 포켓몬' in t or '특수 조우' in t or '이벤트' in t:
                            labels.append(t)
                if not loc_candidates and not notes:
                    continue
                loc = sorted(loc_candidates)[0][1] if loc_candidates else '특수 이벤트'
                lvl = sorted(levels)[0][1] if levels else ''
                method = '특수 이벤트' if any(('조우' in n or '배틀' in n or pokemon_by_id[no]['name'] in n) for n in notes) or '특수 조우' in ' '.join(labels) else '선물'
                add_unique(gifts, {'pokemonId': no, 'pokemonName': pokemon_by_id[no]['name'], 'location': loc, 'method': method, 'level': lvl, 'rates': [], 'source':'GiFTSEVENTS', 'notes': notes[:10]})

    trades = []
    if 'NPCTRADES' in wb.sheetnames:
        ws = wb['NPCTRADES']
        last_loc = ''
        for r in range(1, ws.max_row + 1):
            row_locs = []
            for c in range(1, 12):
                t = clean(ws.cell(r,c).value)
                if looks_location(t):
                    row_locs.append(normalize_location(t))
            if row_locs:
                last_loc = row_locs[-1]
            for c in range(1, ws.max_column + 1):
                t = clean(ws.cell(r,c).value)
                if '받는 교환' not in t: continue
                m = re.search(r'(.+?)를 주고 (.+?)를 받는 교환', t)
                received = m.group(2).strip() if m else ''
                target_id = None
                for p in pokemons:
                    if p['name'] == received or p['name'] in received:
                        target_id = p['id']; break
                if not target_id: continue
                loc = last_loc or 'NPC 교환'
                add_unique(trades, {'pokemonId': target_id, 'pokemonName': pokemon_by_id[target_id]['name'], 'location': loc, 'method': 'NPC 교환', 'level':'', 'rates': [], 'source':'NPCTRADES', 'notes':[t]})

    all_encounters = []
    add_unique.seen = set()
    for e in encounters + gifts + trades:
        # normalize method strings for UI
        e['location'] = normalize_location(e.get('location'))
        e['method'] = clean(e.get('method')) or '상세 조건 확인'
        e['notes'] = [clean(n).rstrip('.') for n in (e.get('notes') or []) if clean(n)]
        add_unique(all_encounters, e)

    for e in all_encounters:
        p = pokemon_by_id.get(e['pokemonId'])
        if p: p['locations'].append(e)
    for p in pokemons:
        # dedupe per pokemon with broader key
        seen = set(); locs=[]
        for e in sorted(p['locations'], key=lambda e: (story_rank(e['location']), e['location'], e['method'], e.get('level',''))):
            key = (e['location'], e['method'], e.get('level',''), tuple(e.get('notes') or []))
            if key not in seen:
                seen.add(key); locs.append(e)
        p['locations'] = locs

    loc_map = {}
    for e in all_encounters:
        loc = e.get('location') or '미분류 지역'
        loc_map.setdefault(loc, {'name': loc, 'order': story_rank(loc), 'encounters': []})['encounters'].append(e)
    locations = sorted(loc_map.values(), key=lambda x: (x['order'], x['name']))

    meta = {
        'sourceWorkbook': '레니게이드_플레티넘_정보_2차_문장번역.xlsx',
        'pokemonCount': len(pokemons), 'encounterCount': len(encounters),
        'giftEventCount': len(gifts) + len(trades), 'locationCount': len(locations),
        'spriteSource': 'DP 앞모습 스프라이트', 'version': 'v7-data-normalized'
    }
    data = {'meta': meta, 'storyOrder': STORY_ORDER, 'pokemons': sorted(pokemons, key=lambda p:p['id']), 'locations': locations, 'encounters': all_encounters}
    OUT.write_text('window.RP_DATA = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
    print(meta)
    for check in [6,238,124,419,418,94]:
        p = pokemon_by_id.get(check)
        if p: print(check, p['name'], p['types'], 'final', p['isFinalEvolution'], 'pre', p.get('preEvolutionIds'), 'locs', len(p['locations']))

if __name__ == '__main__':
    main()
